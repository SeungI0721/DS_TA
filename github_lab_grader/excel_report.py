"""canonical 기록에서 재생성 가능한 로컬 XLSX 보고서를 만든다."""

from __future__ import annotations

import hashlib
import os
import re
import subprocess
import tempfile
from collections import Counter
from collections.abc import Callable, Iterable, Mapping
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .gradebook import SectionGradebook, WeekReportStatus, natural_section_key
from .models import WeeklyRubric
from .record_store import validate_record

_ILLEGAL_XML = re.compile(r"[\x00-\x08\x0B\x0C\x0E-\x1F]")
_FORMULA_PREFIXES = ("=", "+", "-", "@")
_INVALID_SHEET = re.compile(r"[\\/*?:\[\]]")


class ReportError(RuntimeError):
    def __init__(self, code: str, message: str) -> None:
        self.code = code
        super().__init__(message)


def safe_excel_text(value: object | None) -> str | None:
    """외부 문자열에서 수식 실행과 잘못된 XML 제어 문자를 차단한다."""
    if value is None:
        return None
    text = _ILLEGAL_XML.sub("�", str(value))
    return "'" + text if text.startswith(_FORMULA_PREFIXES) else text


def safe_sheet_name(value: str) -> str:
    """Excel 제한을 만족하는 결정적 worksheet 이름을 만든다."""
    cleaned = _INVALID_SHEET.sub("_", value).strip("'") or "Sheet"
    if len(cleaned) <= 31:
        return cleaned
    digest = hashlib.sha256(cleaned.encode("utf-8")).hexdigest()[:6]
    return f"{cleaned[:24]}_{digest}"


def _style_sheet(sheet, header_row: int, last_row: int, last_col: int) -> None:
    sheet.sheet_view.showGridLines = False
    sheet.freeze_panes = f"A{header_row + 1}"
    sheet.auto_filter.ref = f"A{header_row}:{get_column_letter(last_col)}{max(last_row, header_row)}"
    for cell in sheet[f"A{header_row}:{get_column_letter(last_col)}{header_row}"][0]:
        cell.font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F4E78")
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for row in sheet.iter_rows(min_row=header_row + 1, max_row=last_row, max_col=last_col):
        for cell in row:
            cell.font = Font(name="Arial", size=10)
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    for column in range(1, last_col + 1):
        values = [str(sheet.cell(row, column).value or "") for row in range(header_row, last_row + 1)]
        sheet.column_dimensions[get_column_letter(column)].width = min(max(10, max(map(len, values), default=0) + 2), 42)


def _criteria_sheet(workbook: Workbook, records: tuple[dict[str, Any], ...], rubric: WeeklyRubric) -> None:
    sheet = workbook.active
    sheet.title = "채점기준"
    sheet.append(["항목", "값", "설명"])
    rows = [
        ("주차", rubric.week, rubric.title),
        ("최대 점수", rubric.max_score, "canonical 기록과 대조된 주차 배점"),
        ("만점 조건", rubric.score_rules.full, "모든 필수 조건 충족"),
        ("부분 점수 조건", rubric.score_rules.partial, "제출·collaborator 충족, README 신원 일부 누락"),
        ("실패 점수", rubric.score_rules.fail, "필수 제출 또는 collaborator 조건 불충족"),
        ("제출 시각 정책", "ON_TIME / EXTENDED_ON_TIME / LATE", "설정된 네 경계 시각 안의 PushEvent만 사용"),
        ("학생 Push actor 일치", str(rubric.require_student_push_actor), "제출 이벤트 actor와 학생 GitHub ID 일치"),
        ("교수 Collaborator 필수", str(rubric.grading.professor_collaborator_required), "ACTIVE 상태 필요"),
        ("조교 Collaborator 필수", str(rubric.grading.assistant_collaborator_required), "ACTIVE 상태 필요"),
        ("README 필수", str(rubric.grading.readme_required), "선택된 제출 commit에서 확인"),
        ("학번 일치 규칙", "정확한 전체 문자열", "README에서 공식 학번 확인"),
        ("이름 일치 규칙", "Unicode NFC 정확 일치", "README에서 공식 이름 확인"),
        ("수동 검토/null 정책", "점수 셀 공란", "불확실성·오류를 0점으로 변환하지 않음"),
    ]
    for row in rows:
        sheet.append([safe_excel_text(row[0]), row[1], safe_excel_text(row[2])])
    start = sheet.max_row + 2
    sheet.cell(start, 1, "분반")
    timing_headers = ["submission_window_start", "scheduled_deadline", "effective_deadline", "late_window_end", "late_window_source", "deadline_note"]
    for offset, value in enumerate(timing_headers, 2):
        sheet.cell(start, offset, value)
    for record in records:
        timing = record["timing_policy"]
        sheet.append([safe_excel_text(record["section"]), *[safe_excel_text(timing.get(field)) for field in timing_headers]])
    _style_sheet(sheet, 1, sheet.max_row, 7)
    for cell in sheet[start]:
        cell.font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="548235")
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = f"A{start}:G{sheet.max_row}"
    for row in range(start + 1, sheet.max_row + 1):
        sheet.cell(row, 1).number_format = "@"


_WEEKLY_HEADERS = [
    "학번", "이름", "GitHub ID", "Repository", "점수", "배점", "채점 상태", "제출 상태",
    "교수 Collaborator", "조교 Collaborator", "README 상태", "학번 확인", "이름 확인",
    "제출 시각", "제출 Commit SHA", "확인 필요 사유", "오류 코드", "Record Revision",
]


def _weekly_section_sheet(workbook: Workbook, record: dict[str, Any]) -> None:
    sheet = workbook.create_sheet(safe_sheet_name(f"Section{record['section']}"))
    sheet.append(_WEEKLY_HEADERS)
    for student in record["students"]:
        sheet.append([
            safe_excel_text(student["student_id"]), safe_excel_text(student["student_name"]), safe_excel_text(student["github_id"]), safe_excel_text(student["repository"]),
            student["score"], student["max_score"], safe_excel_text(student["grading_status"]), safe_excel_text(student["submission_status"]),
            safe_excel_text(student["professor_collaborator_status"]), safe_excel_text(student["assistant_collaborator_status"]), safe_excel_text(student["readme_status"]),
            safe_excel_text(student["student_id_match"]), safe_excel_text(student["student_name_match"]), safe_excel_text(student.get("submission_push_created_at")),
            safe_excel_text(student.get("submission_push_head_sha")), safe_excel_text(student.get("manual_review_reason")), safe_excel_text(student.get("error_code")), record["record_revision"],
        ])
    _style_sheet(sheet, 1, 1 + len(record["students"]), len(_WEEKLY_HEADERS))
    for row in range(2, 2 + len(record["students"])):
        for column in (1, 3, 4, 15):
            sheet.cell(row, column).number_format = "@"
        for column in (5, 6):
            sheet.cell(row, column).number_format = "0.0"


def _weekly_summary_sheet(workbook: Workbook, records: tuple[dict[str, Any], ...]) -> None:
    sheet = workbook.create_sheet("요약")
    headers = ["분반", "학생 수", "1.0", "0.5", "0.0", "수동 검토", "검증 불가", "오류", "미해결/null", "Record Revision"]
    sheet.append(headers)
    totals = Counter()
    keys = ("total_students", "graded_1_0", "graded_0_5", "graded_0_0", "manual_review", "unverifiable", "errors", "ungraded_or_null")
    for record in records:
        summary = record["summary"]
        values = {key: summary[key] for key in keys}
        totals.update(values)
        sheet.append([safe_excel_text(record["section"]), *values.values(), record["record_revision"]])
    sheet.append(["전체", *[totals[key] for key in keys], None])
    _style_sheet(sheet, 1, sheet.max_row, len(headers))
    for row in range(2, sheet.max_row):
        sheet.cell(row, 1).number_format = "@"


def _weekly_workbook(records: Iterable[Mapping[str, Any]], rubric: WeeklyRubric) -> Workbook:
    checked = tuple(sorted((validate_record(dict(record)) for record in records), key=lambda item: natural_section_key(item["section"])))
    if not checked:
        raise ReportError("REPORT_INPUT_ERROR", "weekly report requires at least one canonical record")
    if any(record["week"] != rubric.week for record in checked):
        raise ReportError("REPORT_INPUT_ERROR", "weekly records and rubric week do not match")
    if len({record["section"] for record in checked}) != len(checked):
        raise ReportError("REPORT_INPUT_ERROR", "weekly report contains a duplicate section")
    if any(record["section"] not in rubric.sections for record in checked):
        raise ReportError("REPORT_INPUT_ERROR", "weekly record section is not configured in the rubric")
    if any(float(record["grading_policy"].get("max_score", -1)) != rubric.max_score for record in checked):
        raise ReportError("REPORT_INPUT_ERROR", "weekly canonical and rubric max scores do not match")
    workbook = Workbook()
    _criteria_sheet(workbook, checked, rubric)
    for record in checked:
        _weekly_section_sheet(workbook, record)
    _weekly_summary_sheet(workbook, checked)
    return workbook


def _completion_label(row) -> str:
    if row.final_weighted_score is not None:
        return "완료"
    statuses = {week.report_status for week in row.weeks}
    if WeekReportStatus.ERROR in statuses:
        return "ERROR"
    if WeekReportStatus.MANUAL_REVIEW in statuses:
        return "MANUAL_REVIEW"
    return "미완료"


def _final_sheet(workbook: Workbook, title: str, gradebooks: tuple[SectionGradebook, ...], *, include_section: bool) -> None:
    expected_weeks = gradebooks[0].expected_weeks if gradebooks else ()
    sheet = workbook.create_sheet(safe_sheet_name(title))
    headers = (["분반"] if include_section else []) + ["학번", "이름", "GitHub ID"] + [f"Week{week:02d}" for week in expected_weeks]
    headers += ["취득점수", "전체 예정 배점", "최종 실습성적", "완료 상태", "확인 필요 수", "오류 수"]
    sheet.append(headers)
    for gradebook in gradebooks:
        for item in gradebook.rows:
            sheet.append((([safe_excel_text(item.section)] if include_section else []) + [
                safe_excel_text(item.student_id), safe_excel_text(item.name), safe_excel_text(item.github_id), *[week.score for week in item.weeks],
                item.earned_score, item.configured_possible_score, item.final_weighted_score, _completion_label(item), item.manual_review_count, item.error_count,
            ]))
    _style_sheet(sheet, 1, sheet.max_row, len(headers))
    id_col = 2 if include_section else 1
    final_col = len(headers) - 3
    for row in range(2, sheet.max_row + 1):
        sheet.cell(row, id_col).number_format = "@"
        sheet.cell(row, id_col + 2).number_format = "@"
        if include_section:
            sheet.cell(row, 1).number_format = "@"
        for column in range(id_col + 3, id_col + 3 + len(expected_weeks)):
            sheet.cell(row, column).number_format = "0.0"
        sheet.cell(row, final_col).number_format = "0.00"


def _weekly_status_sheet(workbook: Workbook, gradebooks: tuple[SectionGradebook, ...]) -> None:
    sheet = workbook.create_sheet("주차별현황")
    headers = ["주차", "분반", "학생 수", "1.0", "0.5", "0.0", "확인 필요", "오류", "Canonical Record Revision", "보고서 생성 상태"]
    sheet.append(headers)
    for book in gradebooks:
        for index, week in enumerate(book.expected_weeks):
            values = [row.weeks[index] for row in book.rows]
            scores = [item.score for item in values]
            revisions = sorted({item.record_revision for item in values if item.record_revision is not None})
            sheet.append([week, safe_excel_text(book.section), len(values), scores.count(1.0), scores.count(0.5), scores.count(0.0),
                          sum(item.report_status is WeekReportStatus.MANUAL_REVIEW for item in values),
                          sum(item.report_status is WeekReportStatus.ERROR for item in values),
                          ",".join(map(str, revisions)) or None,
                          "완료" if values and all(item.score is not None for item in values) else "미완료"])
    _style_sheet(sheet, 1, sheet.max_row, len(headers))
    for row in range(2, sheet.max_row + 1):
        sheet.cell(row, 2).number_format = "@"


class ExcelReportWriter:
    """ignored output/excel 아래 XLSX를 원자적으로 교체한다."""

    def __init__(self, output_root: Path = Path("output/excel"), *, project_root: Path = Path("."), tracked_checker: Callable[[Path], bool] | None = None) -> None:
        self.output_root = output_root.resolve()
        self.project_root = project_root.resolve()
        self._tracked_checker = tracked_checker or self._git_tracked

    def _git_tracked(self, path: Path) -> bool:
        try:
            relative = path.resolve().relative_to(self.project_root)
        except ValueError:
            return False
        try:
            result = subprocess.run(["git", "-C", str(self.project_root), "ls-files", "--error-unmatch", "--", relative.as_posix()], capture_output=True, check=False, timeout=5)
        except (OSError, subprocess.SubprocessError) as exc:
            raise ReportError("OUTPUT_SECURITY_ERROR", "cannot verify output Git privacy") from exc
        return result.returncode == 0

    def _target(self, filename: str) -> Path:
        target = (self.output_root / filename).resolve()
        if not target.is_relative_to(self.output_root):
            raise ReportError("UNSAFE_OUTPUT_PATH", "report path escapes output root")
        if self._tracked_checker(self.output_root) or self._tracked_checker(target):
            raise ReportError("OUTPUT_SECURITY_ERROR", "tracked output path blocks reporting")
        return target

    def weekly_path(self, week: object) -> Path:
        if isinstance(week, bool) or not isinstance(week, int) or not 0 < week <= 999:
            raise ReportError("UNSAFE_OUTPUT_PATH", "week must be a positive integer")
        return self._target(f"week{week:02d}_results.xlsx")

    def final_path(self) -> Path:
        return self._target("final_practical_grade.xlsx")

    def _save(self, workbook: Workbook, target: Path) -> Path:
        target.parent.mkdir(parents=True, exist_ok=True)
        descriptor, temporary_name = tempfile.mkstemp(prefix=f".{target.name}.", suffix=".tmp.xlsx", dir=target.parent)
        os.close(descriptor)
        temporary = Path(temporary_name)
        try:
            workbook.save(temporary)
            workbook.close()
            os.replace(temporary, target)
        except PermissionError as exc:
            raise ReportError("OUTPUT_FILE_LOCKED", f"report file is locked: {target.name}") from exc
        except OSError as exc:
            raise ReportError("REPORT_WRITE_ERROR", f"cannot write report: {target.name}") from exc
        except Exception as exc:
            raise ReportError("REPORT_WRITE_ERROR", f"cannot serialize report: {target.name}") from exc
        finally:
            workbook.close()
            if temporary.exists():
                try:
                    temporary.unlink()
                except OSError:
                    pass
        return target

    def write_weekly_report(self, records: Iterable[Mapping[str, Any]], rubric: WeeklyRubric) -> Path:
        return self._save(_weekly_workbook(records, rubric), self.weekly_path(rubric.week))

    def write_final_report(self, gradebooks: Iterable[SectionGradebook]) -> Path:
        items = tuple(sorted(gradebooks, key=lambda item: natural_section_key(item.section)))
        if not items:
            raise ReportError("REPORT_INPUT_ERROR", "final report requires configured sections")
        expected = items[0].expected_weeks
        if any(item.expected_weeks != expected for item in items):
            raise ReportError("REPORT_INPUT_ERROR", "section gradebooks have different expected weeks")
        workbook = Workbook()
        workbook.remove(workbook.active)
        names: set[str] = set()
        for item in items:
            name = safe_sheet_name(f"Section{item.section}")
            if name.casefold() in names:
                raise ReportError("SHEET_NAME_COLLISION", "section worksheet names collide")
            names.add(name.casefold())
            _final_sheet(workbook, name, (item,), include_section=False)
        _final_sheet(workbook, "전체", items, include_section=True)
        _weekly_status_sheet(workbook, items)
        return self._save(workbook, self.final_path())
