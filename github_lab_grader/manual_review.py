"""Canonical 미확정 결과를 검토용 workbook과 private adjustment로 연결한다."""

from __future__ import annotations

import os
import subprocess
import tempfile
from dataclasses import dataclass
from pathlib import Path
from typing import Any, Callable, Iterable, Mapping

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from .display_labels import format_component_status, format_display_label
from .manual_grade_adjustments import (
    ManualGradeAdjustment,
    allowed_adjustment_scores,
    write_manual_grade_adjustments,
)
from .models import CourseConfig, Student, WeeklyRubric


HEADERS = (
    "분반", "학번", "주차", "실습", "프로젝트 경로", "자동점수", "자동판정",
    "자동판정 사유", "내부 사유 코드", "확인할 내용", "확정된 자동점수 합계",
    "수동확정 점수", "수동조정 사유",
)
_INPUT_SCORE = HEADERS.index("수동확정 점수") + 1
_INPUT_REASON = HEADERS.index("수동조정 사유") + 1


class ManualReviewError(ValueError):
    def __init__(self, code: str, message: str) -> None:
        self.code = code
        super().__init__(f"{format_display_label(code)}: {message}")


@dataclass(frozen=True, slots=True)
class ManualReviewImportResult:
    new: int
    unchanged: int
    skipped: int
    adjustments: tuple[ManualGradeAdjustment, ...]
    written: bool


_HINTS = {
    "PROJECT_MISSING": "마감 시점에 필수 프로젝트 폴더가 실제로 없었는지 확인",
    "COMPUTATION_UNVERIFIABLE": "실습 요구 계산이 코드에서 정상적으로 수행되는지 직접 확인",
    "STUDENT_SOURCE_COMPILE_FAILURE": "학생 코드 자체 오류인지 checker 제약인지 확인",
    "MISSING_REPOSITORY_INFO": "학생의 GitHub 저장소 등록 정보 확인",
    "NOT_FOUND": "저장소 삭제·이름 변경·비공개 전환·재생성 여부 확인",
    "HISTORICAL_EVIDENCE_UNAVAILABLE": "마감 시점의 제출 snapshot을 신뢰성 있게 확인",
    "CHECKER_ERROR": "checker 환경과 제출 소스를 분리해 확인",
    "GRADER_ENVIRONMENT_FAILURE": "사용 가능한 C compiler/runtime 환경에서 다시 확인",
}


def _reason(student: Mapping[str, Any], component: Mapping[str, Any] | None = None) -> str:
    if component:
        return str(component.get("reason") or "WEEK2_MANUAL_REVIEW_REQUIRED")
    return str(
        student.get("error_code") or student.get("manual_review_reason")
        or student.get("submission_status") or "MANUAL_REVIEW"
    )


def extract_manual_review_rows(records: Iterable[Mapping[str, Any]]) -> tuple[dict[str, Any], ...]:
    """점수가 null인 학생 중 실제 검토 대상 component만 펼친다."""
    rows: list[dict[str, Any]] = []
    for record in records:
        section, week = str(record["section"]), int(record["week"])
        for student in record["students"]:
            if student.get("score") is not None:
                continue
            components = [
                item for item in student.get("component_results", ())
                if item.get("status") == "UNVERIFIABLE" or item.get("score") is None
            ]
            targets: tuple[Mapping[str, Any] | None, ...] = tuple(components) or (None,)
            resolved_sum = sum(
                float(item["score"]) for item in student.get("component_results", ())
                if item.get("score") is not None
            )
            for component in targets:
                reason = _reason(student, component)
                status = component.get("status") if component else student.get("grading_status")
                rows.append({
                    "section": section, "student_id": str(student["student_id"]), "week": week,
                    "practice": f"실습{component['practice_number']}" if component else "주차 전체",
                    "project_path": component.get("project_path") if component else None,
                    "automatic_score": component.get("score") if component else student.get("score"),
                    "status": status, "reason": reason,
                    "hint": _HINTS.get(reason, "원본 증거와 rubric을 대조해 최종 주차 점수를 확인"),
                    "resolved_sum": resolved_sum,
                })
    return tuple(rows)


class ManualReviewWorkbook:
    def __init__(
        self, output_root: Path = Path("output/manual_review"), *,
        project_root: Path = Path("."), tracked_checker: Callable[[Path], bool] | None = None,
    ) -> None:
        self.output_root = output_root.resolve()
        self.project_root = project_root.resolve()
        self._tracked_checker = tracked_checker or self._git_tracked

    def _git_tracked(self, path: Path) -> bool:
        try:
            relative = path.resolve().relative_to(self.project_root)
        except ValueError:
            return False
        result = subprocess.run(
            ["git", "-C", str(self.project_root), "ls-files", "--error-unmatch", "--", relative.as_posix()],
            capture_output=True, check=False, timeout=5,
        )
        return result.returncode == 0

    def path(self, week: int) -> Path:
        if isinstance(week, bool) or not isinstance(week, int) or not 0 < week <= 999:
            raise ManualReviewError("UNSAFE_OUTPUT_PATH", "week는 양의 정수여야 합니다")
        target = (self.output_root / f"week{week:02d}_manual_review.xlsx").resolve()
        if not target.is_relative_to(self.output_root) or any(
            self._tracked_checker(item) for item in (self.output_root, target)
        ):
            raise ManualReviewError("OUTPUT_SECURITY_ERROR", "검토 workbook 경로가 Git에 추적됩니다")
        return target

    def write(
        self, records: Iterable[Mapping[str, Any]], rubric: WeeklyRubric
    ) -> Path:
        checked = tuple(records)
        if not checked or any(int(item["week"]) != rubric.week for item in checked):
            raise ManualReviewError("CANONICAL_RECORD_REQUIRED", "해당 주차의 canonical record가 필요합니다")
        rows = extract_manual_review_rows(checked)
        workbook = Workbook()
        guide = workbook.active
        guide.title = "안내"
        guide.append(["수동 검토 workbook 사용법"])
        notes = (
            "이 파일에는 수동 확인이 필요한 학생만 표시됩니다.",
            "자동채점 결과와 내부 식별 열은 수정하지 마세요.",
            "노란색 '수동확정 점수'와 '수동조정 사유'만 입력하세요.",
            "수동확정 점수는 현재 실습 부분점수가 아니라 해당 학생의 최종 주차 총점입니다.",
            "점수를 비워 두면 미결 상태로 유지됩니다.",
            "입력 후 import-manual-review --week XX --dry-run으로 먼저 검증하세요.",
            "Excel은 입력 화면이며 최종 원본은 private manual_grade_adjustments.csv입니다.",
        )
        for note in notes:
            guide.append([note])
        guide.column_dimensions["A"].width = 95
        guide["A1"].font = Font(name="Arial", size=14, bold=True)
        guide.sheet_view.showGridLines = False

        by_section: dict[str, list[dict[str, Any]]] = {}
        for row in rows:
            by_section.setdefault(row["section"], []).append(row)
        sections = tuple(str(item["section"]) for item in checked)
        for section in sections:
            sheet = workbook.create_sheet(f"Section{section}")
            sheet.append(HEADERS)
            for item in by_section.get(section, []):
                sheet.append([
                    item["section"], item["student_id"], item["week"], item["practice"],
                    item["project_path"], item["automatic_score"],
                    format_component_status(item["status"]), format_display_label(item["reason"]),
                    item["reason"], item["hint"], item["resolved_sum"], None, None,
                ])
            _style_review_sheet(sheet, allowed_adjustment_scores(rubric, section))

        summary = workbook.create_sheet("요약")
        summary.append(["분반", "검토 행", "검토 학생"])
        for section in sections:
            section_rows = by_section.get(section, [])
            summary.append([section, len(section_rows), len({item["student_id"] for item in section_rows})])
        _style_review_sheet(summary, ())
        return self._save(workbook, self.path(rubric.week))

    @staticmethod
    def _save(workbook: Workbook, target: Path) -> Path:
        target.parent.mkdir(parents=True, exist_ok=True)
        temporary: Path | None = None
        try:
            descriptor, name = tempfile.mkstemp(prefix=f".{target.name}.", suffix=".tmp.xlsx", dir=target.parent)
            os.close(descriptor)
            temporary = Path(name)
            workbook.save(temporary)
            os.replace(temporary, target)
            return target
        except OSError as exc:
            raise ManualReviewError("REPORT_WRITE_ERROR", "검토 workbook을 저장할 수 없습니다") from exc
        finally:
            workbook.close()
            if temporary is not None:
                temporary.unlink(missing_ok=True)


def _style_review_sheet(sheet, allowed_scores: tuple[float, ...]) -> None:
    sheet.sheet_view.showGridLines = False
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = f"A1:{get_column_letter(sheet.max_column)}{max(1, sheet.max_row)}"
    for cell in sheet[1]:
        cell.font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F4E78")
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.font = Font(name="Arial", size=10)
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    for column in range(1, sheet.max_column + 1):
        width = 16 if column not in (8, 10, 13) else 38
        sheet.column_dimensions[get_column_letter(column)].width = width
    if sheet.max_column >= _INPUT_REASON:
        for column in (_INPUT_SCORE, _INPUT_REASON):
            for row in range(2, sheet.max_row + 1):
                sheet.cell(row, column).fill = PatternFill("solid", fgColor="FFF2CC")
        for row in range(2, sheet.max_row + 1):
            sheet.cell(row, 2).number_format = "@"
            sheet.cell(row, _INPUT_SCORE).number_format = "0.00"
        if allowed_scores and sheet.max_row >= 2:
            values = ",".join(f"{score:g}" for score in allowed_scores)
            validation = DataValidation(type="list", formula1=f'"{values}"', allow_blank=True)
            validation.error = "이 주차에서 허용되는 최종 총점만 입력하세요."
            validation.errorTitle = "허용되지 않는 점수"
            validation.showErrorMessage = True
            sheet.add_data_validation(validation)
            validation.add(f"{get_column_letter(_INPUT_SCORE)}2:{get_column_letter(_INPUT_SCORE)}{sheet.max_row}")


def import_manual_review_workbook(
    workbook_path: Path, *, week: int, course: CourseConfig, students: Iterable[Student],
    rubric: WeeklyRubric, records: Iterable[Mapping[str, Any]],
    existing: Iterable[ManualGradeAdjustment], adjustment_path: Path,
    dry_run: bool = False,
) -> ManualReviewImportResult:
    """Literal 입력만 검증하여 기존 private adjustment와 충돌 없이 병합한다."""
    roster = {(item.section, item.student_id) for item in students}
    eligible = {
        (row["section"], row["week"], row["student_id"])
        for row in extract_manual_review_rows(records)
    }
    try:
        workbook = load_workbook(workbook_path, data_only=False, read_only=False)
    except Exception as exc:
        raise ManualReviewError("MALFORMED_REVIEW_WORKBOOK", "검토 workbook을 읽을 수 없습니다") from exc
    candidates: list[ManualGradeAdjustment] = []
    skipped = 0
    try:
        for sheet in workbook.worksheets:
            if not sheet.title.startswith("Section"):
                continue
            headers = tuple(cell.value for cell in sheet[1])
            if headers != HEADERS:
                raise ManualReviewError("MALFORMED_REVIEW_WORKBOOK", f"{sheet.title} header가 변경되었습니다")
            for values in sheet.iter_rows(min_row=2, values_only=True):
                if not any(value is not None for value in values):
                    continue
                section, student_id, row_week = str(values[0]), str(values[1]), values[2]
                score, reason = values[_INPUT_SCORE - 1], values[_INPUT_REASON - 1]
                if score is None or score == "":
                    skipped += 1
                    continue
                if isinstance(score, bool) or not isinstance(score, (int, float)):
                    code = "FORMULA_INPUT_REJECTED" if isinstance(score, str) and score.startswith("=") else "INVALID_MANUAL_SCORE"
                    raise ManualReviewError(code, f"{section}/{student_id} 점수는 literal 숫자여야 합니다")
                if not isinstance(reason, str) or not reason.strip() or reason.startswith("="):
                    raise ManualReviewError("MANUAL_ADJUSTMENT_REASON_REQUIRED", f"{section}/{student_id} 사유가 필요합니다")
                key = (section, row_week, student_id)
                if row_week != week or week != rubric.week or section not in course.sections:
                    raise ManualReviewError("INVALID_REVIEW_IDENTITY", f"{section}/{student_id} 분반 또는 주차가 잘못되었습니다")
                if (section, student_id) not in roster or key not in eligible:
                    raise ManualReviewError("INVALID_REVIEW_IDENTITY", f"{section}/{student_id}는 현재 검토 대상이 아닙니다")
                if float(score) not in allowed_adjustment_scores(rubric, section):
                    raise ManualReviewError("INVALID_MANUAL_SCORE", f"{section}/{student_id} 점수가 rubric에 맞지 않습니다")
                candidates.append(ManualGradeAdjustment(section, week, student_id, float(score), reason.strip()))
    finally:
        workbook.close()
    keys = [(item.section, item.week, item.student_id) for item in candidates]
    if len(keys) != len(set(keys)):
        raise ManualReviewError("DUPLICATE_REVIEW_DECISION", "같은 학생의 결정이 여러 행에 입력되었습니다")
    merged = {(item.section, item.week, item.student_id): item for item in existing}
    new = unchanged = 0
    for item in candidates:
        key = (item.section, item.week, item.student_id)
        prior = merged.get(key)
        if prior is None:
            merged[key] = item
            new += 1
        elif prior == item:
            unchanged += 1
        else:
            raise ManualReviewError("MANUAL_ADJUSTMENT_CONFLICT", f"{item.section}/{item.student_id} 기존 조정과 충돌합니다")
    result = tuple(sorted(merged.values(), key=lambda item: (item.section, item.week, item.student_id)))
    if not dry_run and new:
        write_manual_grade_adjustments(adjustment_path, result)
    return ManualReviewImportResult(new, unchanged, skipped, result, not dry_run and bool(new))
