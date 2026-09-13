from dataclasses import replace
from pathlib import Path
import subprocess

from openpyxl import load_workbook
import pytest

from github_lab_grader.config_loader import load_weekly_rubric
from github_lab_grader.manual_grade_adjustments import (
    ManualGradeAdjustment,
    load_manual_grade_adjustments,
)
from github_lab_grader.manual_review import (
    HEADERS,
    ManualReviewError,
    ManualReviewWorkbook,
    extract_manual_review_rows,
    import_manual_review_workbook,
)
from github_lab_grader.models import GradingRules, ScoringMode, ScoreRules
from phase6_helpers import make_course, make_record, make_result, make_rubric, make_student


def week2_inputs(tmp_path: Path):
    rubric = load_weekly_rubric(Path(__file__).resolve().parents[1] / "rubrics/week02.json")
    course = make_course(("01",), weeks=2)
    student = make_student("EXAMPLE001")
    record = make_record(course, rubric, results=[make_result(student, week=2, score=None)])
    record["students"][0]["component_results"] = [
        {
            "component_id": "week02-01-practice1", "project_path": "week02-01",
            "practice_number": 1, "status": "PASS", "score": 0.25,
            "max_score": 0.25, "reason": "PASS",
        },
        {
            "component_id": "week02-04-practice7", "project_path": "week02-04",
            "practice_number": 7, "status": "UNVERIFIABLE", "score": None,
            "max_score": 0.25, "reason": "COMPUTATION_UNVERIFIABLE",
        },
    ]
    writer = ManualReviewWorkbook(
        tmp_path / "output/manual_review", project_root=tmp_path,
        tracked_checker=lambda _: False,
    )
    return course, student, rubric, record, writer


def enter_decision(path: Path, score=1.0, reason="fictional TA approval") -> None:
    workbook = load_workbook(path)
    sheet = workbook["Section01"]
    sheet.cell(2, HEADERS.index("수동확정 점수") + 1, score)
    sheet.cell(2, HEADERS.index("수동조정 사유") + 1, reason)
    workbook.save(path)
    workbook.close()


def test_report_contains_only_unresolved_component_and_korean_guidance(tmp_path: Path) -> None:
    course, student, rubric, record, writer = week2_inputs(tmp_path)
    resolved = make_record(course, rubric, results=[make_result(student, week=2, score=1.0)])
    assert extract_manual_review_rows([resolved]) == ()
    path = writer.write([record], rubric)
    workbook = load_workbook(path)
    assert workbook.sheetnames == ["안내", "Section01", "요약"]
    sheet = workbook["Section01"]
    assert sheet.max_row == 2
    values = [cell.value for cell in sheet[2]]
    assert values[1] == "EXAMPLE001" and values[3] == "실습7"
    assert values[4] == "week02-04" and values[5] is None
    assert "수동 확인 필요" in values[6]
    assert "계산 로직" in values[7]
    assert values[8] == "COMPUTATION_UNVERIFIABLE"
    assert values[10] == 0.25
    assert values[11] is None
    assert "최종 주차 총점" in workbook["안내"]["A5"].value
    assert sheet.freeze_panes == "A2" and sheet.auto_filter.ref
    assert len(sheet.data_validations.dataValidation) == 1
    workbook.close()


def test_report_with_no_unresolved_students_has_clear_empty_section(tmp_path: Path) -> None:
    course, student, rubric, _, writer = week2_inputs(tmp_path)
    record = make_record(course, rubric, results=[make_result(student, week=2, score=1.0)])
    path = writer.write([record], rubric)
    workbook = load_workbook(path)
    assert workbook["Section01"].max_row == 1
    assert workbook["요약"]["B2"].value == 0
    workbook.close()


def test_missing_repository_and_not_found_are_review_rows(tmp_path: Path) -> None:
    course, student, rubric, record, _ = week2_inputs(tmp_path)
    first = record["students"][0]
    first["component_results"] = []
    first.update(error_code="MISSING_REPOSITORY_INFO", manual_review_reason=None)
    second = dict(first, student_id="EXAMPLE002", error_code="NOT_FOUND")
    record["students"].append(second)
    rows = extract_manual_review_rows([record])
    assert [item["reason"] for item in rows] == ["MISSING_REPOSITORY_INFO", "NOT_FOUND"]


def test_valid_import_is_atomic_preserves_unrelated_and_dry_run_writes_nothing(tmp_path: Path) -> None:
    course, student, rubric, record, writer = week2_inputs(tmp_path)
    path = writer.write([record], rubric)
    enter_decision(path, 1.0)
    adjustment_path = tmp_path / "data/manual_grade_adjustments.csv"
    unrelated = ManualGradeAdjustment("01", 1, "EXAMPLE001", 1.0, "older fictional approval")
    preview = import_manual_review_workbook(
        path, week=2, course=course, students=[student], rubric=rubric,
        records=[record], existing=[unrelated], adjustment_path=adjustment_path, dry_run=True,
    )
    assert (preview.new, preview.unchanged, preview.skipped, preview.written) == (1, 0, 0, False)
    assert not adjustment_path.exists()
    result = import_manual_review_workbook(
        path, week=2, course=course, students=[student], rubric=rubric,
        records=[record], existing=[unrelated], adjustment_path=adjustment_path,
    )
    assert result.written and adjustment_path.exists()
    loaded = load_manual_grade_adjustments(adjustment_path, course, [student], {1: binary_rubric(), 2: rubric})
    assert {(item.week, item.score) for item in loaded} == {(1, 1.0), (2, 1.0)}


def binary_rubric():
    rubric = make_rubric(1)
    return replace(
        rubric, grading=replace(rubric.grading, scoring_mode=ScoringMode.BINARY),
        score_rules=ScoreRules(1.0, 0.0, 0.0),
    )


def test_week1_binary_import_accepts_only_zero_or_one(tmp_path: Path) -> None:
    rubric = binary_rubric(); course = make_course(("01",), weeks=1); student = make_student()
    record = make_record(course, rubric, results=[make_result(student, score=None)])
    writer = ManualReviewWorkbook(tmp_path / "review", project_root=tmp_path, tracked_checker=lambda _: False)
    path = writer.write([record], rubric)
    enter_decision(path, 1.0)
    result = import_manual_review_workbook(
        path, week=1, course=course, students=[student], rubric=rubric, records=[record],
        existing=(), adjustment_path=tmp_path / "adjustments.csv", dry_run=True,
    )
    assert result.new == 1
    enter_decision(path, 0.5)
    with pytest.raises(ManualReviewError, match="INVALID_MANUAL_SCORE"):
        import_manual_review_workbook(
            path, week=1, course=course, students=[student], rubric=rubric, records=[record],
            existing=(), adjustment_path=tmp_path / "adjustments.csv", dry_run=True,
        )


@pytest.mark.parametrize(
    ("score", "reason", "code"),
    [
        (1.0, "", "MANUAL_ADJUSTMENT_REASON_REQUIRED"),
        (0.3, "fictional", "INVALID_MANUAL_SCORE"),
        ("=1", "fictional", "FORMULA_INPUT_REJECTED"),
    ],
)
def test_import_rejects_invalid_literal_inputs(tmp_path: Path, score, reason, code: str) -> None:
    course, student, rubric, record, writer = week2_inputs(tmp_path)
    path = writer.write([record], rubric); enter_decision(path, score, reason)
    with pytest.raises(ManualReviewError, match=code):
        import_manual_review_workbook(
            path, week=2, course=course, students=[student], rubric=rubric, records=[record],
            existing=(), adjustment_path=tmp_path / "adjustments.csv", dry_run=True,
        )


def test_blank_score_is_ignored_and_conflict_or_duplicate_is_blocked(tmp_path: Path) -> None:
    course, student, rubric, record, writer = week2_inputs(tmp_path)
    path = writer.write([record], rubric)
    result = import_manual_review_workbook(
        path, week=2, course=course, students=[student], rubric=rubric, records=[record],
        existing=(), adjustment_path=tmp_path / "adjustments.csv", dry_run=True,
    )
    assert result.skipped == 1 and result.new == 0
    enter_decision(path, 1.0, "new fictional")
    existing = [ManualGradeAdjustment("01", 2, "EXAMPLE001", 0.75, "old fictional")]
    with pytest.raises(ManualReviewError, match="MANUAL_ADJUSTMENT_CONFLICT"):
        import_manual_review_workbook(
            path, week=2, course=course, students=[student], rubric=rubric, records=[record],
            existing=existing, adjustment_path=tmp_path / "adjustments.csv", dry_run=True,
        )
    workbook = load_workbook(path); sheet = workbook["Section01"]
    # Excel sorting is safe; duplicated decisions for one identity are not.
    values = [cell.value for cell in sheet[2]]
    sheet.append(values)
    workbook.save(path); workbook.close()
    with pytest.raises(ManualReviewError, match="DUPLICATE_REVIEW_DECISION"):
        import_manual_review_workbook(
            path, week=2, course=course, students=[student], rubric=rubric, records=[record],
            existing=(), adjustment_path=tmp_path / "adjustments.csv", dry_run=True,
        )


def test_identical_existing_adjustment_is_idempotent(tmp_path: Path) -> None:
    course, student, rubric, record, writer = week2_inputs(tmp_path)
    path = writer.write([record], rubric); enter_decision(path, 1.0, "fictional approval")
    prior = ManualGradeAdjustment("01", 2, "EXAMPLE001", 1.0, "fictional approval")
    result = import_manual_review_workbook(
        path, week=2, course=course, students=[student], rubric=rubric, records=[record],
        existing=[prior], adjustment_path=tmp_path / "adjustments.csv", dry_run=True,
    )
    assert result.new == 0 and result.unchanged == 1


@pytest.mark.parametrize(("column", "value"), [(1, "02"), (2, "UNKNOWN"), (3, 1)])
def test_changed_identity_fields_are_rejected(tmp_path: Path, column: int, value) -> None:
    course, student, rubric, record, writer = week2_inputs(tmp_path)
    path = writer.write([record], rubric); enter_decision(path)
    workbook = load_workbook(path); workbook["Section01"].cell(2, column, value)
    workbook.save(path); workbook.close()
    with pytest.raises(ManualReviewError, match="INVALID_REVIEW_IDENTITY"):
        import_manual_review_workbook(
            path, week=2, course=course, students=[student], rubric=rubric, records=[record],
            existing=(), adjustment_path=tmp_path / "adjustments.csv", dry_run=True,
        )


def test_private_manual_review_paths_are_git_ignored() -> None:
    root = Path(__file__).resolve().parents[1]
    for relative in ("output/manual_review/week02_manual_review.xlsx", "data/manual_grade_adjustments.csv"):
        result = subprocess.run(["git", "check-ignore", "-q", relative], cwd=root, check=False)
        assert result.returncode == 0
