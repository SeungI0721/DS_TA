"""Phase 6 통합 XLSX 구조, 타입, 보안 및 원자적 출력 테스트."""

from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook

from github_lab_grader.excel_report import ExcelReportWriter, ReportError, safe_excel_text, safe_sheet_name
from github_lab_grader.gradebook import build_all_section_gradebooks
from github_lab_grader.manual_grade_adjustments import (
    ManualGradeAdjustment,
    apply_adjustments_to_weekly_records,
    apply_manual_grade_adjustments,
)
from phase6_helpers import NOW, make_course, make_record, make_result, make_rubric, make_student


@pytest.fixture
def writer(tmp_path: Path) -> ExcelReportWriter:
    return ExcelReportWriter(tmp_path / "output" / "excel", project_root=tmp_path, tracked_checker=lambda _: False)


def weekly_fixture():
    course = make_course(("02", "01"), weeks=1)
    rubric = make_rubric(1, sections=course.sections)
    records = []
    for section in course.sections:
        students = [
            make_result(make_student(f"{section}01", section=section, name="가상학생"), score=1.0),
            make_result(make_student(f"{section}02", section=section, name="부분학생"), score=0.5),
            make_result(make_student(f"{section}03", section=section, name="영점학생"), score=0.0),
            make_result(make_student(f"{section}04", section=section, name="검토학생"), score=None, sha="d" * 40),
        ]
        records.append(make_record(course, rubric, section=section, results=students))
    return course, rubric, records


def test_weekly_workbook_has_criteria_dynamic_sections_summary_and_exact_types(writer: ExcelReportWriter) -> None:
    _, rubric, records = weekly_fixture()
    path = writer.write_weekly_report(records, rubric)
    assert path.name == "week01_results.xlsx"
    workbook = load_workbook(path, data_only=False)
    assert workbook.sheetnames == ["채점기준", "Section01", "Section02", "요약"]
    section = workbook["Section01"]
    assert section["A1"].value == "학번" and section.freeze_panes == "A2"
    assert [section.cell(row, 5).value for row in range(2, 6)] == [1.0, 0.5, 0.0, None]
    assert section["A2"].number_format == "@" and section["S5"].value == "d" * 40
    summary = workbook["요약"]
    assert summary.max_row == 4
    assert [summary.cell(4, column).value for column in range(2, 10)] == [8, 2, 2, 2, 2, 2, 0, 2]
    workbook.close()


def test_weekly_report_rejects_canonical_rubric_max_mismatch(writer: ExcelReportWriter) -> None:
    course = make_course(weeks=1)
    canonical_rubric = make_rubric(1, maximum=1.0)
    current_rubric = make_rubric(1, maximum=2.0)
    record = make_record(course, canonical_rubric)
    with pytest.raises(ReportError, match="max scores do not match"):
        writer.write_weekly_report([record], current_rubric)


def test_weekly_criteria_contains_rubric_and_per_section_timing(writer: ExcelReportWriter) -> None:
    _, rubric, records = weekly_fixture()
    workbook = load_workbook(writer.write_weekly_report(records, rubric), data_only=False)
    values = [[cell.value for cell in row] for row in workbook["채점기준"].iter_rows()]
    flat = " ".join(str(value or "") for row in values for value in row)
    assert all(value in flat for value in ("최대 점수", "수동 검토/null 정책", "scheduled_deadline", "effective_deadline", "late_window_end", "deadline_note"))
    assert flat.count("2026-") >= 8
    assert {row[0] for row in values if row[0] in {"01", "02"}} == {"01", "02"}
    workbook.close()


def test_weekly_report_exposes_required_path_group_results(writer: ExcelReportWriter) -> None:
    course = make_course(weeks=1)
    rubric = make_rubric(1)
    record = make_record(course, rubric)
    record["students"][0]["required_path_checks"] = [
        {
            "name": "repository_root_readme",
            "match": "ANY",
            "paths": ["README.md"],
            "matched_path": "README.md",
            "status": "SATISFIED",
            "failure_reason": "ROOT_README_MISSING",
        },
        {
            "name": "week_project_readme",
            "match": "ANY",
            "paths": ["week01/README.md", "week01-01/README.md"],
            "matched_path": "week01-01/README.md",
            "status": "SATISFIED",
            "failure_reason": "PROJECT_README_MISSING",
        },
    ]
    workbook = load_workbook(
        writer.write_weekly_report([record], rubric), data_only=False
    )
    sheet = workbook["Section01"]
    assert [sheet.cell(1, column).value for column in range(21, 24)] == [
        "Root README",
        "Project README",
        "Project README 경로",
    ]
    assert [sheet.cell(2, column).value for column in range(21, 24)] == [
        "SATISFIED",
        "SATISFIED",
        "week01-01/README.md",
    ]
    workbook.close()


def test_weekly_report_exposes_manual_override_resolution(writer: ExcelReportWriter) -> None:
    course = make_course(weeks=1); rubric = make_rubric(1)
    record = make_record(course, rubric)
    record["students"][0].update(
        {
            "grade_resolution_source": "MANUAL_OVERRIDE",
            "manual_override_action": "PRESERVE_PREVIOUS_CANONICAL_RESULT",
            "manual_override_reason": "fictional approved exception",
            "manual_override_source_revision": 1,
            "manual_override_source_run_id": record["grading_run_id"],
        }
    )
    workbook = load_workbook(writer.write_weekly_report([record], rubric), data_only=False)
    sheet = workbook["Section01"]
    assert sheet["X1"].value == "Resolution Source"
    assert sheet["Y1"].value == "Override Reason"
    assert sheet["X2"].value == "MANUAL_OVERRIDE"
    assert sheet["Y2"].value == "fictional approved exception"
    workbook.close()


def test_weekly_report_separates_automatic_adjustment_and_effective_scores(writer: ExcelReportWriter) -> None:
    course = make_course(weeks=1); rubric = make_rubric(1)
    record = make_record(course, rubric, results=[make_result(make_student(), score=None)])
    projected = apply_adjustments_to_weekly_records(
        [record],
        (ManualGradeAdjustment("01", 1, "EXAMPLE001", 1.0, "=fictional approval"),),
    )
    workbook = load_workbook(writer.write_weekly_report(projected, rubric), data_only=False)
    sheet = workbook["Section01"]
    assert [sheet.cell(1, column).value for column in range(5, 10)] == [
        "자동점수", "수동조정", "최종점수", "점수출처", "수동조정사유",
    ]
    assert [sheet.cell(2, column).value for column in range(5, 9)] == [
        None, 1.0, 1.0, "MANUAL_ADJUSTMENT",
    ]
    assert sheet["I2"].data_type != "f" and str(sheet["I2"].value).startswith("'")
    workbook.close()


def test_final_report_has_dynamic_sections_overall_status_and_completion_semantics(writer: ExcelReportWriter) -> None:
    course = make_course(("10", "02", "01"), weeks=2, weight=10.0)
    rubrics = {1: make_rubric(1, sections=course.sections, maximum=2.0), 2: make_rubric(2, sections=course.sections, maximum=1.0)}
    roster = [make_student(f"ID{s}", section=s) for s in course.sections]
    records = []
    for index, section in enumerate(course.sections):
        records.append(make_record(course, rubrics[1], section=section, results=[make_result(roster[index], score=0.0, maximum=2.0)]))
        if section != "02":
            records.append(make_record(course, rubrics[2], section=section, results=[make_result(roster[index], week=2, score=1.0)]))
    books = build_all_section_gradebooks(course, records, rubrics, roster, as_of=NOW)
    path = writer.write_final_report(books)
    workbook = load_workbook(path, data_only=False)
    assert path.name == "final_practical_grade.xlsx"
    assert workbook.sheetnames == ["Section01", "Section02", "Section10", "전체", "주차별현황"]
    complete = workbook["Section01"]
    assert complete["D2"].value == 0.0 and complete["E2"].value == 1.0
    assert complete["H2"].value == pytest.approx(10 / 3) and complete["I2"].value == "완료"
    incomplete = workbook["Section02"]
    assert incomplete["E2"].value is None and incomplete["H2"].value is None and incomplete["I2"].value == "미완료"
    assert [workbook["전체"].cell(row, 1).value for row in range(2, 5)] == ["01", "02", "10"]
    assert workbook["전체"]["A2"].number_format == "@" and workbook["전체"]["D2"].number_format == "@"
    assert workbook["주차별현황"].max_row == 7
    workbook.close()


def test_final_report_uses_effective_manual_adjustment_score(writer: ExcelReportWriter) -> None:
    course = make_course(("01",), weeks=1, weight=10.0)
    rubric = make_rubric(1, sections=course.sections)
    student = make_student("EXAMPLE001", section="01")
    records = [make_record(course, rubric, section="01", results=[make_result(student, score=None)])]
    books = apply_manual_grade_adjustments(
        build_all_section_gradebooks(course, records, {1: rubric}, [student], as_of=NOW),
        [ManualGradeAdjustment("01", 1, "EXAMPLE001", 1.0, "Verified by instructor")],
    )
    workbook = load_workbook(writer.write_final_report(books), data_only=False)
    sheet = workbook["Section01"]
    assert (sheet["D2"].value, sheet["E2"].value, sheet["G2"].value) == (1.0, 1.0, 10.0)
    workbook.close()


@pytest.mark.parametrize("payload", ["=1+1", "+CMD", "-1+2", "@SUM(A1:A2)"])
def test_safe_excel_text_neutralizes_formula_prefixes(payload: str) -> None:
    assert safe_excel_text(payload) == "'" + payload


def test_untrusted_weekly_fields_are_never_formulas(writer: ExcelReportWriter) -> None:
    course = make_course(weeks=1); rubric = make_rubric(1)
    malicious = make_student("0001", name="=1+1", github_id="+CMD", repository="-1+2")
    record = make_record(course, rubric, results=[make_result(malicious, score=None, reason="@SUM(A1:A2)")])
    workbook = load_workbook(writer.write_weekly_report([record], rubric), data_only=False)
    sheet = workbook["Section01"]
    for coordinate in ("B2", "C2", "D2", "Z2"):
        assert sheet[coordinate].data_type != "f" and str(sheet[coordinate].value).startswith("'")
    workbook.close()


def test_control_characters_are_safely_replaced() -> None:
    assert safe_excel_text("가상\x01학생") == "가상�학생"


def test_safe_sheet_name_is_deterministic_and_bounded() -> None:
    value = safe_sheet_name("Section" + "a" * 40)
    assert value == safe_sheet_name("Section" + "a" * 40) and len(value) <= 31
    assert safe_sheet_name("A/B:C") == "A_B_C"


@pytest.mark.parametrize("week", [0, -1, "../1", True])
def test_output_path_rejects_invalid_week(writer: ExcelReportWriter, week) -> None:
    with pytest.raises(ReportError):
        writer.weekly_path(week)


def test_tracked_output_blocks_report(tmp_path: Path) -> None:
    writer = ExcelReportWriter(tmp_path / "output" / "excel", tracked_checker=lambda _: True)
    with pytest.raises(ReportError, match="tracked output"):
        writer.weekly_path(1)


def test_replace_failure_preserves_old_report_and_cleans_temp(writer: ExcelReportWriter, monkeypatch: pytest.MonkeyPatch) -> None:
    _, rubric, records = weekly_fixture(); target = writer.write_weekly_report(records, rubric); old = target.read_bytes()
    monkeypatch.setattr("github_lab_grader.excel_report.os.replace", lambda *_: (_ for _ in ()).throw(PermissionError("locked")))
    with pytest.raises(ReportError) as caught:
        writer.write_weekly_report(records, rubric)
    assert caught.value.code == "OUTPUT_FILE_LOCKED" and target.read_bytes() == old
    assert not list(target.parent.glob("*.tmp.xlsx"))


def test_workbook_save_failure_preserves_old_report(writer: ExcelReportWriter, monkeypatch: pytest.MonkeyPatch) -> None:
    _, rubric, records = weekly_fixture(); target = writer.write_weekly_report(records, rubric); old = target.read_bytes()
    monkeypatch.setattr(Workbook, "save", lambda *_: (_ for _ in ()).throw(OSError("simulated")))
    with pytest.raises(ReportError) as caught:
        writer.write_weekly_report(records, rubric)
    assert caught.value.code == "REPORT_WRITE_ERROR" and target.read_bytes() == old
    assert not list(target.parent.glob("*.tmp.xlsx"))


def test_workbook_generation_does_not_use_network(writer: ExcelReportWriter, monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.setattr("socket.create_connection", lambda *_args, **_kwargs: pytest.fail("network access"))
    course, rubric, records = weekly_fixture()
    books = build_all_section_gradebooks(course, records, {1: rubric})
    paths = [writer.write_weekly_report(records, rubric), writer.write_final_report(books)]
    for path in paths:
        workbook = load_workbook(path); workbook.close()
